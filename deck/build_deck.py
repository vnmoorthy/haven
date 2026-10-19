#!/usr/bin/env python3
"""Build deck/Haven-Pitch.pptx — the Haven investor / hackathon pitch deck.

Re-runnable:  python3 deck/build_deck.py
  * Embeds real screenshots from docs/screenshots/*.png when present
    (mission-control.png, incident-room.png, audit.png); otherwise draws a
    stylized mock UI so the deck never shows a broken image.
  * Adds Morph transitions (fade fallback) to every slide and staggered
    entrance animations (fade / rise / zoom / motion path) via <p:timing>.

Requires: python-pptx>=0.6.21, Pillow, lxml.
"""
import io
import math
import os
from itertools import count

from lxml import etree
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SHOTS = os.environ.get("HAVEN_SHOTS", os.path.join(ROOT, "docs", "screenshots"))
ASSETS = os.path.join(HERE, "assets")
OUT = os.environ.get("HAVEN_DECK_OUT", os.path.join(HERE, "Haven-Pitch.pptx"))
os.makedirs(ASSETS, exist_ok=True)

FONT = "Arial"  # brand font is Inter; Arial ships everywhere and keeps widths predictable
BG = "05060A"
INK = "E6E8EF"
MUTED = "8B90A0"
VIOLET = "8B5CF6"
VIOLET_L = "A78BFA"
CYAN = "22D3EE"
ROSE = "F43F5E"
EMERALD = "10B981"
AMBER = "F59E0B"
WHITE = "FFFFFF"
PANEL = "FFFFFF"

SW, SH = 13.333, 7.5
GITHUB = "https://github.com/vnmoorthy/haven"


def hexrgb(h):
    return RGBColor.from_string(h.upper())


def rgb_t(h):
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


# ----------------------------------------------------------------------------
# Raster assets (Pillow)
# ----------------------------------------------------------------------------
def _font(size, bold=False):
    cands = (["/System/Library/Fonts/Supplemental/Arial Bold.ttf", "/Library/Fonts/Arial Bold.ttf",
              "C:/Windows/Fonts/arialbd.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"]
             if bold else
             ["/System/Library/Fonts/Supplemental/Arial.ttf", "/Library/Fonts/Arial.ttf",
              "C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"])
    for c in cands:
        if os.path.exists(c):
            return ImageFont.truetype(c, size)
    return ImageFont.load_default()


def make_orb(path, color, size=900):
    """Soft radial glow with alpha falloff."""
    img = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    px = img.load()
    r, g, b = rgb_t(color)
    c = size / 2
    for y in range(size):
        for x in range(size):
            d = math.hypot(x - c, y - c) / c
            if d < 1:
                a = (1 - d) ** 2.2
                px[x, y] = (r, g, b, int(255 * 0.55 * a))
    img.save(path)


def make_bg(path):
    """Full-slide background: near-black with a faint dot grid and vignette."""
    W, H = 1920, 1080
    img = Image.new("RGB", (W, H), rgb_t(BG))
    d = ImageDraw.Draw(img)
    for y in range(0, H, 36):
        for x in range(0, W, 36):
            d.point((x, y), fill=(28, 31, 44))
    vign = Image.new("L", (W, H), 0)
    vd = ImageDraw.Draw(vign)
    vd.ellipse((-W * 0.25, -H * 0.35, W * 1.25, H * 1.35), fill=255)
    vign = vign.filter(ImageFilter.GaussianBlur(160))
    dark = Image.new("RGB", (W, H), (2, 3, 6))
    img = Image.composite(img, dark, vign)
    img.save(path)


def _shield_points(s):
    pts = [(32 * s, 4 * s), (8 * s, 13 * s), (8 * s, 30 * s)]

    def bez(p0, p1, p2, p3, n=40):
        out = []
        for i in range(1, n + 1):
            t = i / n
            x = ((1 - t) ** 3) * p0[0] + 3 * ((1 - t) ** 2) * t * p1[0] + 3 * (1 - t) * t * t * p2[0] + t ** 3 * p3[0]
            y = ((1 - t) ** 3) * p0[1] + 3 * ((1 - t) ** 2) * t * p1[1] + 3 * (1 - t) * t * t * p2[1] + t ** 3 * p3[1]
            out.append((x * s, y * s))
        return out

    pts += bez((8, 30), (8, 45), (18.4, 56.7), (32, 60))
    pts += bez((32, 60), (45.6, 56.7), (56, 45), (56, 30))
    pts += [(56 * s, 13 * s)]
    return pts


def make_logo(path, px=512):
    S = px * 2
    s = S / 64
    grad = Image.new("RGBA", (S, S))
    gp = grad.load()
    a, b = rgb_t("A78BFA"), rgb_t("22D3EE")
    for y in range(S):
        for x in range(S):
            t = (x + y) / (2 * S)
            gp[x, y] = tuple(int(a[i] + (b[i] - a[i]) * t) for i in range(3)) + (255,)
    mask = Image.new("L", (S, S), 0)
    ImageDraw.Draw(mask).polygon(_shield_points(s), fill=255)
    img = Image.new("RGBA", (S, S), (0, 0, 0, 0))
    img.paste(grad, (0, 0), mask)
    d = ImageDraw.Draw(img)
    pts = [(22 * s, 33.5 * s), (29 * s, 40 * s), (42.5 * s, 25 * s)]
    w = int(5 * s)
    d.line(pts, fill=(11, 13, 20, 255), width=w, joint="curve")
    for p in (pts[0], pts[-1]):
        d.ellipse((p[0] - w / 2, p[1] - w / 2, p[0] + w / 2, p[1] + w / 2), fill=(11, 13, 20, 255))
    img = img.resize((px, px), Image.LANCZOS)
    img.save(path)


def make_icon(path, kind, px=256):
    S = px * 2
    img = Image.new("RGBA", (S, S), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    W = (255, 255, 255, 255)
    lw = int(S * 0.075)
    c = S / 2
    if kind == "bolt":
        d.polygon([(S * .56, S * .12), (S * .26, S * .56), (S * .48, S * .56), (S * .42, S * .88),
                   (S * .74, S * .42), (S * .52, S * .42)], fill=W)
    elif kind == "search":
        d.ellipse((S * .18, S * .18, S * .64, S * .64), outline=W, width=lw)
        d.line((S * .58, S * .58, S * .84, S * .84), fill=W, width=int(lw * 1.3))
    elif kind == "heart":
        r = S * .19
        d.ellipse((S * .5 - 2 * r, S * .22, S * .5, S * .22 + 2 * r), fill=W)
        d.ellipse((S * .5, S * .22, S * .5 + 2 * r, S * .22 + 2 * r), fill=W)
        d.polygon([(S * .5 - 2 * r + S * .012, S * .22 + r * 1.25), (S * .5 + 2 * r - S * .012, S * .22 + r * 1.25),
                   (S * .5, S * .84)], fill=W)
    elif kind == "user":
        d.ellipse((S * .36, S * .14, S * .64, S * .42), fill=W)
        d.pieslice((S * .18, S * .48, S * .82, S * 1.12), 180, 360, fill=W)
        d.rectangle((0, S * .82, S, S), fill=(0, 0, 0, 0))
    elif kind == "eye":
        d.ellipse((S * .12, S * .3, S * .88, S * .7), outline=W, width=lw)
        d.ellipse((S * .4, S * .4, S * .6, S * .6), fill=W)
        d.line((S * .2, S * .84, S * .8, S * .16), fill=W, width=lw)
    elif kind == "doc":
        d.rounded_rectangle((S * .24, S * .12, S * .76, S * .88), radius=S * .06, outline=W, width=lw)
        for i, yy in enumerate((.32, .46, .6)):
            d.line((S * .34, S * yy, S * (.66 if i < 2 else .52), S * yy), fill=W, width=int(lw * .8))
        d.line((S * .5, S * .74, S * .57, S * .8), fill=W, width=lw)
        d.line((S * .57, S * .8, S * .7, S * .66), fill=W, width=lw)
    elif kind == "lock":
        d.rounded_rectangle((S * .22, S * .44, S * .78, S * .86), radius=S * .06, fill=W)
        d.arc((S * .32, S * .14, S * .68, S * .62), 180, 360, fill=W, width=lw)
        d.line((S * .32, S * .38, S * .32, S * .46), fill=W, width=lw)
        d.line((S * .68 - lw / 2, S * .38, S * .68 - lw / 2, S * .46), fill=W, width=lw)
    elif kind == "lifebuoy":
        d.ellipse((S * .14, S * .14, S * .86, S * .86), outline=W, width=int(lw * 1.6))
        d.ellipse((S * .33, S * .33, S * .67, S * .67), outline=W, width=lw)
    img = img.resize((px, px), Image.LANCZOS)
    img.save(path)


def _mock_base(title, url):
    W, H = 1600, 1000
    img = Image.new("RGB", (W, H), (11, 13, 20))
    d = ImageDraw.Draw(img)
    # sidebar
    d.rectangle((0, 0, 250, H), fill=(15, 17, 26))
    logo = Image.open(os.path.join(ASSETS, "logo.png")).resize((44, 44), Image.LANCZOS)
    img.paste(logo, (28, 28), logo)
    d.text((86, 36), "Haven", font=_font(28, True), fill=(230, 232, 239))
    nav = [("Mission Control", VIOLET), ("Incidents", ROSE), ("Community", CYAN), ("Bot Audits", AMBER),
           ("Ask Haven", EMERALD), ("Team", VIOLET_L)]
    for i, (n, col) in enumerate(nav):
        y = 120 + i * 58
        if n == title:
            d.rounded_rectangle((16, y - 10, 234, y + 38), radius=12, fill=(30, 27, 52))
        d.ellipse((32, y + 8, 44, y + 20), fill=rgb_t(col))
        d.text((58, y + 2), n, font=_font(22), fill=(200, 204, 216) if n == title else (139, 144, 160))
    d.text((290, 34), title, font=_font(40, True), fill=(240, 242, 248))
    d.rounded_rectangle((W - 190, 36, W - 40, 80), radius=22, fill=(8, 46, 38))
    d.ellipse((W - 172, 52, W - 160, 64), fill=rgb_t(EMERALD))
    d.text((W - 150, 45), "LIVE", font=_font(22, True), fill=rgb_t(EMERALD))
    return img, d


def make_mock(path, kind):
    if kind == "mission-control":
        img, d = _mock_base("Mission Control", "haven/mission-control")
        kpis = [("Messages screened", "12,480", CYAN), ("Incidents today", "37", ROSE),
                ("Median response", "0.4s", EMERALD), ("Members protected", "40k", VIOLET_L)]
        for i, (lab, val, col) in enumerate(kpis):
            x = 290 + i * 322
            d.rounded_rectangle((x, 120, x + 300, 290), radius=20, fill=(20, 22, 34), outline=(40, 44, 62), width=2)
            d.text((x + 24, 146), lab, font=_font(22), fill=(139, 144, 160))
            d.text((x + 24, 190), val, font=_font(60, True), fill=rgb_t(col))
        d.text((290, 330), "Live incident feed", font=_font(28, True), fill=(230, 232, 239))
        rows = [(ROSE, "92", "Phishing link quarantined in #general", "Nova -> Rex"),
                (AMBER, "71", "Crisis message - Mira replied in 3s", "Mira"),
                (VIOLET_L, "F", "Bot audit: RaidShield Pro graded F", "Kai"),
                (CYAN, "64", "Raid burst throttled in #welcome", "Nova"),
                (EMERALD, "OK", "Domain added to shared blocklist", "Juno")]
        for i, (col, sc, txt, who) in enumerate(rows):
            y = 380 + i * 112
            d.rounded_rectangle((290, y, 1560, y + 96), radius=18, fill=(17, 19, 30), outline=(36, 40, 56), width=2)
            d.rounded_rectangle((316, y + 22, 400, y + 74), radius=14, fill=rgb_t(col))
            tw = d.textlength(sc, font=_font(28, True))
            d.text((358 - tw / 2, y + 32), sc, font=_font(28, True), fill=(11, 13, 20))
            d.text((428, y + 32), txt, font=_font(28), fill=(220, 224, 234))
            d.text((1360, y + 34), who, font=_font(24), fill=(139, 144, 160))
    elif kind == "incident-room":
        img, d = _mock_base("Incidents", "haven/incidents/1042")
        d.text((290, 100), "Incident #1042  ·  Phishing  ·  risk 92", font=_font(26), fill=(139, 144, 160))
        msgs = [(ROSE, "scammer", "FREE Premium for everyone -> disc0rd-gift.co/claim", "QUARANTINED 0.4s"),
                (CYAN, "Nova", "Risk 92 - lookalike domain, urgency bait. Quarantined.", ""),
                (VIOLET_L, "Aria", "Delegating to Rex over BAND for detonation.", ""),
                (AMBER, "Rex", "Detonated in Celesto sandbox: credential harvester.", ""),
                (EMERALD, "Juno", "disc0rd-gift.co added to the shared blocklist.", "")]
        for i, (col, who, txt, tag) in enumerate(msgs):
            y = 160 + i * 150
            d.ellipse((290, y, 350, y + 60), fill=rgb_t(col))
            d.text((310, y + 14), who[0].upper(), font=_font(28, True), fill=(11, 13, 20))
            d.text((370, y), who, font=_font(24, True), fill=rgb_t(col))
            d.rounded_rectangle((370, y + 36, 1010, y + 120), radius=16, fill=(20, 22, 34))
            d.text((392, y + 50), txt[:44], font=_font(22), fill=(220, 224, 234))
            if len(txt) > 44:
                d.text((392, y + 82), txt[44:], font=_font(22), fill=(220, 224, 234))
            if tag:
                d.rounded_rectangle((520, y - 4, 760, y + 30), radius=14, fill=(70, 14, 28))
                d.text((536, y + 1), tag, font=_font(20, True), fill=rgb_t(ROSE))
        # sandbox screenshot card
        d.rounded_rectangle((1050, 130, 1570, 680), radius=20, fill=(20, 22, 34), outline=(56, 48, 96), width=2)
        d.text((1074, 150), "Celesto sandbox capture", font=_font(22, True), fill=(200, 204, 216))
        d.rounded_rectangle((1074, 196, 1546, 560), radius=14, fill=(236, 238, 245))
        d.text((1180, 230), "Sign in to claim", font=_font(30, True), fill=(30, 32, 44))
        for j in range(2):
            d.rounded_rectangle((1120, 300 + j * 76, 1500, 352 + j * 76), radius=10, fill=(255, 255, 255),
                                outline=(190, 194, 206), width=2)
        d.rounded_rectangle((1120, 460, 1500, 512), radius=10, fill=(88, 101, 242))
        d.text((1074, 584), "Harvests credentials -> 3 redirects", font=_font(22), fill=rgb_t(ROSE))
        d.rounded_rectangle((1050, 720, 1570, 810), radius=22, fill=rgb_t(VIOLET))
        d.text((1192, 744), "Approve ban  (1 click)", font=_font(30, True), fill=(255, 255, 255))
        d.text((1050, 840), "n8n will log enforcement", font=_font(22), fill=(139, 144, 160))
    else:  # audit
        img, d = _mock_base("Bot Audits", "haven/audits/raidshield-pro")
        d.text((290, 100), "RaidShield Pro  ·  third-party bot", font=_font(26), fill=(139, 144, 160))
        d.ellipse((300, 170, 600, 470), fill=(60, 12, 26), outline=rgb_t(ROSE), width=8)
        tw = d.textlength("F", font=_font(200, True))
        d.text((450 - tw / 2, 205), "F", font=_font(200, True), fill=rgb_t(ROSE))
        d.text((300, 500), "Do not install", font=_font(34, True), fill=rgb_t(ROSE))
        finds = [(ROSE, "Snyk", "Known vulnerable dependencies (high severity)"),
                 (ROSE, "Permissions", "Requests ADMINISTRATOR"),
                 (ROSE, "Code", "eval() on raw user input"),
                 (AMBER, "Tenki", "Sandbox run: outbound calls to unknown host"),
                 (AMBER, "Prelint", "Violates bot-policy spec (2 rules)")]
        for i, (col, tag, txt) in enumerate(finds):
            y = 170 + i * 128
            d.rounded_rectangle((680, y, 1560, y + 108), radius=18, fill=(17, 19, 30), outline=(36, 40, 56), width=2)
            d.rounded_rectangle((704, y + 30, 880, y + 78), radius=14, fill=rgb_t(col))
            tw = d.textlength(tag, font=_font(24, True))
            d.text((792 - tw / 2, y + 40), tag, font=_font(24, True), fill=(11, 13, 20))
            d.text((904, y + 38), txt, font=_font(26), fill=(220, 224, 234))
    img.save(path)


def build_assets():
    paths = {}
    p = os.path.join(ASSETS, "logo.png")
    if not os.path.exists(p):
        make_logo(p)
    paths["logo"] = p
    p = os.path.join(ASSETS, "bg.png")
    if not os.path.exists(p):
        make_bg(p)
    paths["bg"] = p
    for name, col in (("orb_violet", VIOLET), ("orb_cyan", CYAN), ("orb_rose", ROSE)):
        p = os.path.join(ASSETS, name + ".png")
        if not os.path.exists(p):
            make_orb(p, col)
        paths[name] = p
    for k in ("bolt", "search", "heart", "user", "eye", "doc", "lock", "lifebuoy"):
        p = os.path.join(ASSETS, "icon_%s.png" % k)
        if not os.path.exists(p):
            make_icon(p, k)
        paths["icon_" + k] = p
    for kind in ("mission-control", "incident-room", "audit"):
        p = os.path.join(ASSETS, "mock_%s.png" % kind)
        make_mock(p, kind)
        paths["mock_" + kind] = p
    return paths


# ----------------------------------------------------------------------------
# Shape helpers (python-pptx + raw XML)
# ----------------------------------------------------------------------------
A = None  # asset paths, set in main()
_name_seq = count(1)


def _uname(prefix="el"):
    return "%s_%d" % (prefix, next(_name_seq))


def _alpha(clr_parent, alpha):
    clr = clr_parent[0]
    for old in clr.findall(qn("a:alpha")):
        clr.remove(old)
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(alpha * 100000)))


def _drop_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def rect(shapes, x, y, w, h, fill=PANEL, alpha=0.05, line=WHITE, line_alpha=0.10, line_w=0.75,
         radius=0.12, shape=None, name=None, grad=None, glow=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    s = shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    _drop_style(s)
    s.name = name or _uname("shape")
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE and radius:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    if grad:
        s.fill.gradient()
        s.fill.gradient_angle = grad[2] if len(grad) > 2 else 0
        st = s.fill.gradient_stops
        st[0].color.rgb = hexrgb(grad[0]); st[0].position = 0.0
        st[1].color.rgb = hexrgb(grad[1]); st[1].position = 1.0
    elif fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = hexrgb(fill)
        if alpha is not None and alpha < 1:
            _alpha(s._element.spPr.find(qn("a:solidFill")), alpha)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = hexrgb(line)
        s.line.width = Pt(line_w)
        if line_alpha is not None and line_alpha < 1:
            _alpha(s._element.spPr.find(qn("a:ln")).find(qn("a:solidFill")), line_alpha)
    if glow:
        eff = etree.SubElement(s._element.spPr, qn("a:effectLst"))
        g = etree.SubElement(eff, qn("a:glow"))
        g.set("rad", str(int(Pt(glow[1]))))
        c = etree.SubElement(g, qn("a:srgbClr")); c.set("val", glow[0])
        a = etree.SubElement(c, qn("a:alpha")); a.set("val", str(int(glow[2] * 100000)))
    s.text_frame.text = ""
    return s


def text(shapes, x, y, w, h, runs, size=16, color=INK, bold=False, align="l", anchor="t",
         name=None, spacing=None, line_spacing=None, para_after=None, italic=False, grad=None, link=None):
    """runs: str | list of paragraphs; a paragraph is str or list of (text, opts) tuples."""
    tb = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name or _uname("text")
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if line_spacing:
            p.line_spacing = line_spacing
        if para_after is not None:
            p.space_after = Pt(para_after)
        segs = para if isinstance(para, list) else [(para, {})]
        for seg, o in segs:
            r = p.add_run()
            r.text = seg
            f = r.font
            f.name = o.get("font", FONT)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            g = o.get("grad", grad)
            if g:
                rPr = r._r.get_or_add_rPr()
                gf = etree.SubElement(rPr, qn("a:gradFill"))
                gs = etree.SubElement(gf, qn("a:gsLst"))
                for pos, col in ((0, g[0]), (100000, g[1])):
                    e = etree.SubElement(gs, qn("a:gs")); e.set("pos", str(pos))
                    c = etree.SubElement(e, qn("a:srgbClr")); c.set("val", col)
                lin = etree.SubElement(gf, qn("a:lin")); lin.set("ang", "0"); lin.set("scaled", "0")
                # latin font must come after fill in rPr
                lat = rPr.find(qn("a:latin"))
                if lat is not None:
                    rPr.remove(lat); rPr.append(lat)
            else:
                f.color.rgb = hexrgb(o.get("color", color))
            sp = o.get("spacing", spacing)
            if sp:
                r._r.get_or_add_rPr().set("spc", str(int(sp)))
    if link:
        tb.click_action.hyperlink.address = link
    return tb


def picture(shapes, path, x, y, w=None, h=None, name=None):
    p = shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    p.name = name or _uname("pic")
    return p


def picture_cover(shapes, path, x, y, w, h, name=None):
    """Place an image cropped to fill (cover) the box, anchored to the top."""
    with Image.open(path) as im:
        iw, ih = im.size
    p = shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    p.name = name or _uname("shot")
    box, img = w / h, iw / ih
    if img > box:  # too wide -> crop the right edge (keeps the app's sidebar/nav visible)
        p.crop_right = 1 - box / img
    elif img < box:  # too tall -> crop bottom
        p.crop_bottom = 1 - img / box
    return p


def line(shapes, x1, y1, x2, y2, color=VIOLET_L, width=1.25, alpha=0.6, arrow=False, dash=False, name=None):
    c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    _drop_style(c)
    c.name = name or _uname("line")
    c.line.color.rgb = hexrgb(color)
    c.line.width = Pt(width)
    ln = c._element.spPr.find(qn("a:ln"))
    _alpha(ln.find(qn("a:solidFill")), alpha)
    if dash:
        pd = etree.SubElement(ln, qn("a:prstDash")); pd.set("val", "dash")
    if arrow:
        te = etree.SubElement(ln, qn("a:tailEnd")); te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    return c


def group(shapes, name=None):
    g = shapes.add_group_shape()
    g.name = name or _uname("group")
    return g


def icon_circle(shapes, x, y, d, color, icon, alpha=0.18):
    rect(shapes, x, y, d, d, fill=color, alpha=alpha, line=color, line_alpha=0.55, radius=None,
         shape=MSO_SHAPE.OVAL)
    pad = d * 0.25
    picture(shapes, A["icon_" + icon], x + pad, y + pad, d - 2 * pad, d - 2 * pad)


def pill(shapes, x, y, w, h, label, color=INK, fill=PANEL, alpha=0.06, line=WHITE, line_alpha=0.14, size=14,
         bold=True, dot=None):
    g = group(shapes)
    rect(g.shapes, x, y, w, h, fill=fill, alpha=alpha, line=line, line_alpha=line_alpha, radius=h / 2)
    tx = x
    if dot:
        dd = h * 0.3
        rect(g.shapes, x + h * 0.38, y + (h - dd) / 2, dd, dd, fill=dot, alpha=1, line=None, radius=None,
             shape=MSO_SHAPE.OVAL)
        tx = x + h * 0.38 + dd
    text(g.shapes, tx, y, w - (tx - x), h, label, size=size, color=color, bold=bold, align="c", anchor="m")
    return g


# ----------------------------------------------------------------------------
# Motion: transitions + entrance animations
# ----------------------------------------------------------------------------
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"


def add_morph_transition(slide):
    xml = (
        '<mc:AlternateContent xmlns:mc="%s" xmlns:p="%s">'
        '<mc:Choice xmlns:p159="%s" Requires="p159">'
        '<p:transition spd="slow"><p159:morph option="byObject"/></p:transition>'
        '</mc:Choice>'
        '<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback>'
        '</mc:AlternateContent>' % (NS_MC, NS_P, NS_P159)
    )
    slide._element.append(etree.fromstring(xml))


def _tgt(spid):
    return '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl>' % spid


def _effect_xml(ids, spid, kind, delay, dur, path=None):
    """Return a <p:par> for one 'with previous' entrance effect."""
    eid = next(ids)
    vis = ('<p:set><p:cBhvr><p:cTn id="%d" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
           '</p:cTn>%s<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
           '<p:to><p:strVal val="visible"/></p:to></p:set>')

    def anim(attr, frm, to):
        return ('<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="%d" dur="%d" fill="hold"/>'
                '%s<p:attrNameLst><p:attrName>%s</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst>'
                '<p:tav tm="0"><p:val><p:strVal val="%s"/></p:val></p:tav>'
                '<p:tav tm="100000"><p:val><p:strVal val="%s"/></p:val></p:tav></p:tavLst></p:anim>'
                % (next(ids), dur, _tgt(spid), attr, frm, to))

    def fade():
        return ('<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="%d" dur="%d"/>%s</p:cBhvr>'
                '</p:animEffect>' % (next(ids), dur, _tgt(spid)))

    if kind == "path":
        body = ('<p:animMotion origin="layout" path="%s" pathEditMode="relative" ptsTypes="">'
                '<p:cBhvr><p:cTn id="%d" dur="%d" fill="hold" repeatCount="indefinite"/>%s'
                '<p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst>'
                '</p:cBhvr></p:animMotion>' % (path, next(ids), dur, _tgt(spid)))
        preset = 'presetID="0" presetClass="path" presetSubtype="0"'
    else:
        body = vis % (next(ids), _tgt(spid))
        if kind == "fade":
            preset = 'presetID="10" presetClass="entr" presetSubtype="0"'
            body += fade()
        elif kind == "rise":  # "Rise up": fade + slide up a little
            preset = 'presetID="42" presetClass="entr" presetSubtype="0"'
            body += fade() + anim("ppt_x", "#ppt_x", "#ppt_x") + anim("ppt_y", "#ppt_y+.06", "#ppt_y")
        elif kind == "left":  # fade + drift in from the left
            preset = 'presetID="47" presetClass="entr" presetSubtype="0"'
            body += fade() + anim("ppt_x", "#ppt_x-.06", "#ppt_x") + anim("ppt_y", "#ppt_y", "#ppt_y")
        elif kind == "zoom":
            preset = 'presetID="53" presetClass="entr" presetSubtype="16"'
            body += anim("ppt_w", "0", "#ppt_w") + anim("ppt_h", "0", "#ppt_h") + fade()
        else:
            raise ValueError(kind)
    return ('<p:par><p:cTn id="%d" %s fill="hold" grpId="0" nodeType="withEffect">'
            '<p:stCondLst><p:cond delay="%d"/></p:stCondLst><p:childTnLst>%s</p:childTnLst></p:cTn></p:par>'
            % (eid, preset, delay, body))


def add_animations(slide, effects):
    """effects: list of (shape, kind, delay_ms[, dur_ms[, path]]). Plays automatically after the transition."""
    if not effects:
        return
    ids = count(5)
    pars = []
    bld = []
    for e in effects:
        shp, kind, delay = e[0], e[1], e[2]
        dur = e[3] if len(e) > 3 else 600
        path = e[4] if len(e) > 4 else None
        spid = shp.shape_id
        pars.append(_effect_xml(ids, spid, kind, delay, dur, path))
        if shp._element.tag == qn("p:sp") and kind != "path":
            bld.append('<p:bldP spid="%d" grpId="0" animBg="1"/>' % spid)
    xml = (
        '<p:timing xmlns:p="%s"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
        '<p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
        '<p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
        '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
        '<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        '%s'
        '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
        '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        '</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>%s</p:timing>'
        % (NS_P, "".join(pars), ("<p:bldLst>%s</p:bldLst>" % "".join(bld)) if bld else "")
    )
    slide._element.append(etree.fromstring(xml))


# ----------------------------------------------------------------------------
# Slide scaffolding
# ----------------------------------------------------------------------------
ORBS = {  # per-slide orb positions (x, y, size) — Morph glides them between slides
    1: [(-1.5, -2.0, 8.0), (7.5, 2.5, 8.0), (4.5, 5.2, 5.0)],
    2: [(8.5, -3.0, 7.5), (-3.0, 3.5, 7.0), (10.5, 4.5, 4.5)],
    3: [(-2.5, -3.0, 7.0), (8.0, 3.8, 7.0), (11.0, -1.5, 4.0)],
    4: [(3.0, -4.5, 8.0), (-3.0, 4.0, 6.5), (9.5, 4.0, 5.5)],
    5: [(3.2, 0.2, 7.0), (-3.5, -2.5, 6.0), (10.0, 4.5, 4.5)],
    6: [(9.0, -3.5, 7.0), (-2.5, 3.8, 7.0), (5.0, 5.5, 4.0)],
    7: [(-3.0, -2.0, 7.5), (9.5, 3.0, 6.5), (6.0, -2.5, 4.0)],
    8: [(4.0, -4.0, 7.0), (-3.5, 3.0, 7.0), (10.0, 3.5, 5.0)],
    9: [(9.0, -3.0, 7.0), (-2.5, 3.5, 6.5), (4.5, 5.0, 4.0)],
    10: [(-1.0, -1.5, 8.5), (6.5, 2.0, 8.5), (3.5, 5.0, 5.0)],
}


def new_slide(prs, n, kicker=None, title=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = hexrgb(BG)
    picture(s.shapes, A["bg"], 0, 0, SW, SH, name="!!bg")
    for (x, y, d), key in zip(ORBS[n], ("orb_violet", "orb_cyan", "orb_rose")):
        picture(s.shapes, A[key], x, y, d, d, name="!!" + key)
    # footer
    if n != 1:
        picture(s.shapes, A["logo"], 0.6, 6.92, 0.26, 0.26, name="!!footlogo")
        text(s.shapes, 0.95, 6.92, 3, 0.26, "Haven", size=11, color=MUTED, bold=True, anchor="m", name="!!footname")
        text(s.shapes, SW - 0.6 - 2, 6.92, 2, 0.26, "%02d / 10" % n, size=11, color=MUTED, align="r", anchor="m",
             name="!!pageno")
    t_shapes = []
    if kicker:
        t_shapes.append(text(s.shapes, 0.6, 0.5, 8, 0.3, kicker.upper(), size=12, color=CYAN, bold=True,
                             spacing=300))
    if title:
        t_shapes.append(text(s.shapes, 0.6, 0.82, 12.1, 0.75, title, size=34, color=WHITE, bold=True))
    return s, t_shapes


def screenshot_frame(shapes, key, x, y, w, h, url):
    """Glass browser window with the real screenshot if present, otherwise a stylized mock."""
    g = group(shapes)
    rect(g.shapes, x, y, w, h, fill="0B0D14", alpha=0.92, line=VIOLET_L, line_alpha=0.35, line_w=1, radius=0.14,
         glow=(VIOLET, 18, 0.22))
    bar = 0.32
    for i, col in enumerate((ROSE, AMBER, EMERALD)):
        rect(g.shapes, x + 0.18 + i * 0.17, y + 0.11, 0.1, 0.1, fill=col, alpha=0.9, line=None, radius=None,
             shape=MSO_SHAPE.OVAL)
    rect(g.shapes, x + w / 2 - 1.3, y + 0.07, 2.6, 0.18, fill=WHITE, alpha=0.06, line=None, radius=0.09)
    text(g.shapes, x + w / 2 - 1.3, y + 0.07, 2.6, 0.18, url, size=8, color=MUTED, align="c", anchor="m")
    real = os.path.join(SHOTS, key + ".png")
    src = real if os.path.exists(real) else A["mock_" + key]
    picture_cover(g.shapes, src, x + 0.08, y + bar, w - 0.16, h - bar - 0.08)
    return g


# ----------------------------------------------------------------------------
# Slides
# ----------------------------------------------------------------------------
def slide1(prs):
    s, _ = new_slide(prs, 1)
    fx = []
    logo = picture(s.shapes, A["logo"], SW / 2 - 0.75, 0.95, 1.5, 1.5, name="!!herologo")
    fx.append((logo, "zoom", 0, 800))
    t = text(s.shapes, 1, 2.6, SW - 2, 1.4, "Haven", size=96, bold=True, align="c", anchor="m",
             grad=(VIOLET_L, CYAN), spacing=-200)
    fx.append((t, "rise", 300, 800))
    tag = text(s.shapes, 1.5, 4.05, SW - 3, 0.6, "The AI Trust & Safety team for every online community",
               size=26, color=INK, align="c", anchor="m")
    fx.append((tag, "rise", 700, 700))
    chips = [("8 AI employees", VIOLET_L), ("1 human founder", CYAN), ("14 sponsor tools", EMERALD)]
    cw, gap = 2.7, 0.3
    x0 = SW / 2 - (3 * cw + 2 * gap) / 2
    for i, (lab, col) in enumerate(chips):
        p = pill(s.shapes, x0 + i * (cw + gap), 5.05, cw, 0.56, lab, color=INK, dot=col, size=15)
        fx.append((p, "rise", 1100 + i * 180, 600))
    f = text(s.shapes, 1, 6.55, SW - 2, 0.35, "github.com/vnmoorthy/haven  ·  Founder: Moorthy (CEO)",
             size=12, color=MUTED, align="c", anchor="m", link=GITHUB)
    fx.append((f, "fade", 1800, 600))
    add_animations(s, fx)
    s.notes_slide.notes_text_frame.text = (
        "Hi, I'm Moorthy. This is Haven: the AI Trust and Safety team for every online community. "
        "Haven is an AI-native company: one human founder and eight AI employees who screen messages, "
        "investigate threats, and care for people in crisis, around the clock. We built it on fourteen "
        "sponsor tools. In the next few minutes I'll show you the problem, the team, and a live takedown.")
    return s


def slide2(prs):
    s, ts = new_slide(prs, 2, "The problem", "Volunteers can't keep up with 3 AM threats")
    fx = [(ts[0], "fade", 0), (ts[1], "rise", 150)]
    sub = text(s.shapes, 0.6, 1.62, 11.5, 0.4,
               "Most of the world's 30M+ online communities are moderated by unpaid volunteers. "
               "Harm lands in seconds; human response takes hours.", size=16, color=MUTED)
    fx.append((sub, "fade", 350))
    stats = [("30M+", "online communities on Discord, Slack, Telegram and forums", VIOLET_L),
             ("$15B+", "trust & safety market, growing as platforms push safety down to communities", CYAN),
             ("$4.20", "average human moderator cost to handle one incident", AMBER),
             ("3 AM", "when the scams, raids and crises arrive and nobody is online", ROSE)]
    cw, gap, y = (12.13 - 3 * 0.3) / 4, 0.3, 2.35
    for i, (num, lab, col) in enumerate(stats):
        x = 0.6 + i * (cw + gap)
        g = group(s.shapes)
        rect(g.shapes, x, y, cw, 2.15, radius=0.18)
        text(g.shapes, x + 0.3, y + 0.28, cw - 0.6, 0.85, num, size=44, bold=True, color=col)
        text(g.shapes, x + 0.3, y + 1.15, cw - 0.6, 0.85, lab, size=13, color=MUTED)
        fx.append((g, "rise", 500 + i * 160))
    threats = [("Phishing & fake \u201cfree Premium\u201d links", ROSE),
               ("Raids & spam floods", AMBER),
               ("Malicious third-party bots", VIOLET_L),
               ("Members in crisis, alone at night", CYAN)]
    y2 = 4.85
    lab = text(s.shapes, 0.6, y2, 6, 0.3, "WHAT VOLUNTEERS FACE EVERY NIGHT", size=11, color=MUTED, bold=True,
               spacing=200)
    fx.append((lab, "fade", 1150))
    for i, (t_, col) in enumerate(threats):
        x = 0.6 + i * (cw + gap)
        p = pill(s.shapes, x, y2 + 0.42, cw, 0.62, t_, color=INK, dot=col, size=13, bold=False,
                 fill=col, alpha=0.08, line=col, line_alpha=0.35)
        fx.append((p, "rise", 1250 + i * 120))
    q = text(s.shapes, 0.6, 6.15, 12.1, 0.45,
             [[("Result: ", {"bold": True, "color": INK}),
               ("moderators burn out, scams spread, and the people who most need help wait until morning.",
                {"color": MUTED})]], size=15)
    fx.append((q, "fade", 1800))
    add_animations(s, fx)
    s.notes_slide.notes_text_frame.text = (
        "There are more than thirty million online communities, and almost all of them are protected by "
        "unpaid volunteers. Scams, phishing links, raids and people in crisis show up at 3 AM, when nobody "
        "is online. A human moderator costs about four dollars twenty per incident, so most communities "
        "simply go unprotected. Moderators burn out, and harm spreads in seconds.")
    return s


def slide3(prs):
    s, ts = new_slide(prs, 3, "The solution", "An AI-native T&S team, on call 24/7")
    fx = [(ts[0], "fade", 0), (ts[1], "rise", 150)]
    one = text(s.shapes, 0.6, 1.7, 5.9, 0.95,
               "Haven plugs into your community and runs a full Trust & Safety team of AI employees, "
               "with a human founder approving every ban.", size=16, color=MUTED, line_spacing=1.1)
    fx.append((one, "fade", 350))
    pillars = [("bolt", CYAN, "Detect in under 1 second",
                "Nova screens every message through a RocketRide triage pipeline and quarantines threats instantly."),
               ("search", VIOLET_L, "Investigate with evidence",
                "Rex detonates links in an isolated Celesto computer via AdaL: screenshots, redirects, lookalikes."),
               ("heart", ROSE, "Care for humans",
                "Mira answers crisis messages in seconds with 988 resources, a real face and voice, and pages a human.")]
    for i, (ic, col, h, d) in enumerate(pillars):
        y = 2.85 + i * 1.3
        g = group(s.shapes)
        rect(g.shapes, 0.6, y, 5.9, 1.12, radius=0.16)
        icon_circle(g.shapes, 0.82, y + 0.24, 0.64, col, ic)
        text(g.shapes, 1.7, y + 0.16, 4.6, 0.32, h, size=16, bold=True, color=WHITE)
        text(g.shapes, 1.7, y + 0.5, 4.65, 0.55, d, size=12, color=MUTED)
        fx.append((g, "left", 550 + i * 200))
    fr = screenshot_frame(s.shapes, "mission-control", 6.85, 1.75, 5.88, 4.85, "haven / mission-control")
    fx.append((fr, "zoom", 900, 800))
    add_animations(s, fx)
    s.notes_slide.notes_text_frame.text = (
        "Haven is an AI-native Trust and Safety company. It plugs into your community and does three things. "
        "It detects threats in under a second: Nova screens every message. It investigates with evidence: "
        "Rex detonates suspicious links in a sandbox. And it cares for humans: Mira responds to people in "
        "crisis within seconds. This is Mission Control, where the founder sees everything live.")
    return s


TEAM = [
    ("Nova", "T&S Triage Analyst", "RocketRide · PromptQL", CYAN),
    ("Rex", "Threat Intel Investigator", "AdaL · Celesto · Finch", ROSE),
    ("Mira", "Community Wellbeing Lead", "Masky.AI · n8n", "EC4899"),
    ("Kai", "Security Engineer", "Snyk · Tenki · Prelint", AMBER),
    ("Iris", "Data Analyst", "PromptQL · Paritok", "3B82F6"),
    ("Juno", "Policy & Product Manager", "Prelint · Finch", EMERALD),
    ("Leo", "Growth & Customer Success", "n8n · AdaL · Masky", VIOLET_L),
]


def slide4(prs):
    s, ts = new_slide(prs, 4, "The team", "Meet the team: 1 human, 8 AI employees")
    fx = [(ts[0], "fade", 0), (ts[1], "rise", 150)]
    cx = SW / 2
    # founder
    gf = group(s.shapes)
    rect(gf.shapes, cx - 1.75, 1.75, 3.5, 0.78, fill=WHITE, alpha=0.08, line=WHITE, line_alpha=0.3, radius=0.39)
    picture(gf.shapes, A["icon_user"], cx - 1.55, 1.92, 0.44, 0.44)
    text(gf.shapes, cx - 1.0, 1.82, 2.6, 0.34, "Moorthy", size=16, bold=True, color=WHITE, anchor="m")
    text(gf.shapes, cx - 1.0, 2.15, 2.6, 0.28, "Founder & CEO (human)", size=11, color=MUTED, anchor="m")
    fx.append((gf, "zoom", 300))
    # aria
    ga = group(s.shapes)
    rect(ga.shapes, cx - 2.6, 2.95, 5.2, 0.95, grad=("4C1D95", "0E7490", 0), line=VIOLET_L, line_alpha=0.6,
         radius=0.2, glow=(VIOLET, 14, 0.35))
    rect(ga.shapes, cx - 2.4, 3.12, 0.6, 0.6, fill=WHITE, alpha=0.95, line=None, radius=None, shape=MSO_SHAPE.OVAL)
    text(ga.shapes, cx - 2.4, 3.12, 0.6, 0.6, "A", size=20, bold=True, color="4C1D95", align="c", anchor="m")
    text(ga.shapes, cx - 1.65, 3.05, 4.1, 0.36, [[("Aria", {"bold": True, "color": WHITE, "size": 17}),
                                                   ("   Chief of Staff", {"color": "E9E5FF", "size": 13})]],
         anchor="m")
    text(ga.shapes, cx - 1.65, 3.42, 4.15, 0.4, "Routes every task over BAND, writes the founder briefing",
         size=11, color="D7DCF5", anchor="m")
    fx.append((ga, "zoom", 550))
    l1 = line(s.shapes, cx, 2.53, cx, 2.95, color=VIOLET_L, width=1.5, alpha=0.8)
    fx.append((l1, "fade", 450))
    # specialists
    n = len(TEAM)
    gap = 0.16
    cw = (12.13 - (n - 1) * gap) / n
    y = 4.45
    bus_y = 4.18
    lb = line(s.shapes, 0.6 + cw / 2, bus_y, 0.6 + cw / 2 + (n - 1) * (cw + gap), bus_y, color=VIOLET_L, width=1.25,
              alpha=0.55)
    l2 = line(s.shapes, cx, 3.9, cx, bus_y, color=VIOLET_L, width=1.5, alpha=0.8)
    fx += [(l2, "fade", 750), (lb, "fade", 800)]
    for i, (nm, title, tools, col) in enumerate(TEAM):
        x = 0.6 + i * (cw + gap)
        mx = x + cw / 2
        ld = line(s.shapes, mx, bus_y, mx, y, color=col, width=1.25, alpha=0.7)
        g = group(s.shapes)
        rect(g.shapes, x, y, cw, 2.2, fill=col, alpha=0.07, line=col, line_alpha=0.35, radius=0.16)
        rect(g.shapes, mx - 0.3, y + 0.2, 0.6, 0.6, fill=col, alpha=0.95, line=None, radius=None,
             shape=MSO_SHAPE.OVAL)
        text(g.shapes, mx - 0.3, y + 0.2, 0.6, 0.6, nm[0], size=20, bold=True, color="0B0D14", align="c",
             anchor="m")
        text(g.shapes, x + 0.08, y + 0.9, cw - 0.16, 0.3, nm, size=15, bold=True, color=WHITE, align="c")
        text(g.shapes, x + 0.08, y + 1.2, cw - 0.16, 0.48, title, size=10.5, color=INK, align="c")
        text(g.shapes, x + 0.08, y + 1.68, cw - 0.16, 0.45, tools, size=9.5, color=col, align="c", bold=True)
        fx += [(ld, "fade", 900 + i * 110), (g, "rise", 950 + i * 110)]
    add_animations(s, fx)
    s.notes_slide.notes_text_frame.text = (
        "Here's the company. I'm the only human. Aria, our Chief of Staff, routes every task over BAND and "
        "writes my daily briefing. Under her are seven specialists: Nova triages, Rex investigates links, "
        "Mira handles wellbeing and crisis, Kai audits third-party bots, Iris answers data questions, Juno "
        "owns policy, and Leo runs onboarding and customer success. Each one uses specific sponsor tools.")
    return s


def slide5(prs):
    s, ts = new_slide(prs, 5, "How it works", "How it works: one operating loop")
    fx = [(ts[0], "fade", 0), (ts[1], "rise", 150)]
    nodes = [  # (label, sub, x, y, w, color)
        ("Community", "Discord · Slack · Telegram", 0.6, 1.95, 2.45, INK),
        ("Nova", "RocketRide triage <1ms", 3.45, 1.95, 2.45, CYAN),
        ("BAND", "agent message bus", 6.3, 1.95, 2.0, VIOLET_L),
        ("Aria", "routes & delegates", 8.7, 1.95, 2.0, VIOLET_L),
        # specialists block drawn separately at right
        ("Founder approval", "1-click ban", 8.7, 5.55, 2.0, EMERALD),
        ("n8n", "enforcement log", 6.3, 5.55, 2.0, AMBER),
        ("PromptQL", "learning loop", 3.45, 5.55, 2.45, "3B82F6"),
        ("Finch · Paritok", "skills · compression", 0.6, 5.55, 2.45, "EC4899"),
    ]
    nh = 0.85
    shapes_ = []
    for lab, sub, x, y, w, col in nodes:
        g = group(s.shapes)
        rect(g.shapes, x, y, w, nh, fill=col, alpha=0.1, line=col, line_alpha=0.55, radius=0.2)
        text(g.shapes, x + 0.1, y + 0.1, w - 0.2, 0.34, lab, size=15, bold=True, color=WHITE, align="c")
        text(g.shapes, x + 0.1, y + 0.46, w - 0.2, 0.3, sub, size=11, color=MUTED, align="c")
        shapes_.append(g)
    # specialists block
    sx, sy, sw, sh = 11.05, 1.95, 1.68, 4.45
    gs = group(s.shapes)
    rect(gs.shapes, sx, sy, sw, sh, fill=VIOLET, alpha=0.08, line=VIOLET_L, line_alpha=0.45, radius=0.2)
    text(gs.shapes, sx, sy + 0.12, sw, 0.3, "Specialists", size=13, bold=True, color=WHITE, align="c")
    specs = [("Rex", "AdaL + Celesto", ROSE), ("Mira", "Masky + n8n", "EC4899"), ("Juno", "Prelint", EMERALD),
             ("Kai", "Snyk + Tenki", AMBER)]
    for j, (nm, tl, col) in enumerate(specs):
        yy = sy + 0.55 + j * 0.95
        rect(gs.shapes, sx + 0.12, yy, sw - 0.24, 0.82, fill=col, alpha=0.12, line=col, line_alpha=0.5, radius=0.14)
        text(gs.shapes, sx + 0.12, yy + 0.1, sw - 0.24, 0.3, nm, size=14, bold=True, color=WHITE, align="c")
        text(gs.shapes, sx + 0.12, yy + 0.43, sw - 0.24, 0.3, tl, size=10, color=col, align="c", bold=True)
    # arrows around the loop (clockwise)
    ay_top = 1.95 + nh / 2
    ay_bot = 5.55 + nh / 2
    arrows = [
        line(s.shapes, 3.05, ay_top, 3.45, ay_top, arrow=True, color=CYAN, alpha=0.9),
        line(s.shapes, 5.9, ay_top, 6.3, ay_top, arrow=True, color=CYAN, alpha=0.9),
        line(s.shapes, 8.3, ay_top, 8.7, ay_top, arrow=True, color=CYAN, alpha=0.9),
        line(s.shapes, 10.7, ay_top, 11.05, ay_top, arrow=True, color=CYAN, alpha=0.9),
        line(s.shapes, 11.05, ay_bot, 10.7, ay_bot, arrow=True, color=CYAN, alpha=0.9),
        line(s.shapes, 8.7, ay_bot, 8.3, ay_bot, arrow=True, color=CYAN, alpha=0.9),
        line(s.shapes, 6.3, ay_bot, 5.9, ay_bot, arrow=True, color=CYAN, alpha=0.9),
        line(s.shapes, 3.45, ay_bot, 3.05, ay_bot, arrow=True, color=CYAN, alpha=0.9),
        line(s.shapes, 1.825, 5.55, 1.825, 1.95 + nh, arrow=True, color=CYAN, alpha=0.9, dash=True),
    ]
    order = shapes_[:4] + [gs] + shapes_[4:]
    arrow_after = [0, 1, 2, 3, 4, 5, 6, 7, 8]
    for i, g in enumerate(order):
        fx.append((g, "zoom", 400 + i * 230, 450))
        if i < len(arrows):
            fx.append((arrows[i], "fade", 400 + i * 230 + 180, 300))
    # center panel: Flatkey
    gc = group(s.shapes)
    rect(gc.shapes, 3.45, 3.25, 7.25, 1.85, fill=WHITE, alpha=0.04, line=WHITE, line_alpha=0.1, radius=0.2)
    text(gc.shapes, 3.75, 3.4, 6.7, 0.3, "EVERY CALL GOES THROUGH FLATKEY", size=11, bold=True, color=CYAN,
         spacing=200, align="c")
    tiers = [("fast", CYAN), ("reasoning", VIOLET_L), ("coding", AMBER), ("vision", ROSE)]
    tw_, tg = 1.45, 0.2
    tx0 = 3.45 + (7.25 - (4 * tw_ + 3 * tg)) / 2
    for k, (tn, col) in enumerate(tiers):
        pill(gc.shapes, tx0 + k * (tw_ + tg), 3.82, tw_, 0.42, tn, color=col, fill=col, alpha=0.1, line=col,
             line_alpha=0.5, size=12)
    text(gc.shapes, 3.75, 4.4, 6.7, 0.55,
         "Tiered models with automatic failover and per-call cost tracking. "
         "Kylon holds the workforce manifest; Finch skills are shared by every agent.",
         size=12, color=MUTED, align="c")
    fx.append((gc, "fade", 400 + 9 * 230))
    # travelling dot: a motion path around the loop
    d = 0.2
    rest_y = (1.95 + nh + 5.55) / 2
    dot = rect(s.shapes, 1.825 - d / 2, rest_y - d / 2, d, d, fill=CYAN, alpha=1, line=None, radius=None,
               shape=MSO_SHAPE.OVAL, glow=(CYAN, 10, 0.6), name="loopdot")

    def rel(xi, yi):
        return "%.4f %.4f" % ((xi - 1.825) / SW, (yi - rest_y) / SH)

    pts = [(1.825, rest_y), (1.825, ay_top), (sx + sw / 2, ay_top), (sx + sw / 2, ay_bot), (1.825, ay_bot),
           (1.825, rest_y)]
    path = "M 0 0 " + " ".join("L " + rel(*p) for p in pts[1:]) + " E"
    fx.append((dot, "path", 2900, 9000, path))
    add_animations(s, fx)
    s.notes_slide.notes_text_frame.text = (
        "Every message flows through one loop. Nova screens it with a RocketRide pipeline. Anything risky "
        "goes over BAND to Aria, who delegates to the right specialist: Rex for links, Mira for crisis, Juno "
        "for policy, Kai for bots. Bans come to me for a one-click approval, n8n logs enforcement, and PromptQL "
        "plus Finch skills turn every incident into learning. All model calls run through Flatkey.")
    return s


def slide6(prs):
    s, ts = new_slide(prs, 6, "Live demo", "Live demo: a phishing takedown in 6 seconds")
    fx = [(ts[0], "fade", 0), (ts[1], "rise", 150)]
    steps = [("0.0s", "Posted", "Scammer drops a fake \u201cfree Premium\u201d link", ROSE),
             ("0.4s", "Quarantined", "Nova scores risk 80+ and hides it", AMBER),
             ("2s", "Delegated", "Aria hands it to Rex over BAND", VIOLET_L),
             ("4s", "Detonated", "Celesto sandbox captures the credential harvester", CYAN),
             ("6s", "Warned", "Rex warns the channel by voice; Juno blocklists the domain", "EC4899"),
             ("1 click", "Banned", "Founder approves; n8n logs the enforcement", EMERALD)]
    n = len(steps)
    x0, x1 = 0.6, 12.73
    seg = (x1 - x0) / n
    ty = 2.15
    track = rect(s.shapes, x0 + seg / 2, ty + 0.17, (n - 1) * seg, 0.05, grad=(ROSE, EMERALD, 0), line=None, radius=None)
    fx.append((track, "left", 300, 900))
    for i, (t_, h, d, col) in enumerate(steps):
        cx = x0 + seg * i + seg / 2
        g = group(s.shapes)
        rect(g.shapes, cx - 0.2, ty, 0.4, 0.4, fill=col, alpha=1, line=WHITE, line_alpha=0.7, line_w=1.5,
             radius=None, shape=MSO_SHAPE.OVAL, glow=(col, 10, 0.45))
        text(g.shapes, cx - seg / 2 + 0.05, ty + 0.55, seg - 0.1, 0.42, t_, size=22, bold=True, color=col, align="c")
        text(g.shapes, cx - seg / 2 + 0.05, ty + 0.98, seg - 0.1, 0.3, h, size=14, bold=True, color=WHITE, align="c")
        text(g.shapes, cx - seg / 2 + 0.1, ty + 1.3, seg - 0.2, 0.62, d, size=11, color=MUTED, align="c")
        fx.append((g, "zoom", 600 + i * 330, 450))
    # bottom row: screenshot + 3 moment cards
    by, bh = 4.3, 2.4
    fr = screenshot_frame(s.shapes, "incident-room", 0.6, by, 3.9, bh, "haven / incidents / 1042")
    fx.append((fr, "rise", 2700))
    cards_x = 4.8
    cw = (12.73 - cards_x - 2 * 0.25) / 3
    # crisis
    g1 = group(s.shapes)
    x = cards_x
    rect(g1.shapes, x, by, cw, bh, fill="EC4899", alpha=0.07, line="EC4899", line_alpha=0.35, radius=0.16)
    icon_circle(g1.shapes, x + 0.22, by + 0.22, 0.5, "EC4899", "lifebuoy")
    text(g1.shapes, x + 0.86, by + 0.3, cw - 1.0, 0.36, "Crisis care", size=15, bold=True, color=WHITE)
    text(g1.shapes, x + 0.22, by + 0.84, cw - 0.44, 0.48, "< 5s", size=30, bold=True, color="F472B6")
    text(g1.shapes, x + 0.22, by + 1.3, cw - 0.44, 0.25, "to a caring reply", size=11, color=MUTED)
    text(g1.shapes, x + 0.22, by + 1.62, cw - 0.44, 0.7,
         "\u201cI want to end it\u201d \u2192 Mira answers with 988 resources; n8n pages a human.",
         size=11, color=INK)
    # audit
    g2 = group(s.shapes)
    x = cards_x + cw + 0.25
    rect(g2.shapes, x, by, cw, bh, fill=ROSE, alpha=0.07, line=ROSE, line_alpha=0.35, radius=0.16)
    rect(g2.shapes, x + 0.22, by + 0.22, 0.82, 0.82, fill=ROSE, alpha=0.18, line=ROSE, line_alpha=0.8, radius=None,
         shape=MSO_SHAPE.OVAL)
    text(g2.shapes, x + 0.22, by + 0.22, 0.82, 0.82, "F", size=34, bold=True, color=ROSE, align="c", anchor="m")
    text(g2.shapes, x + 1.18, by + 0.3, cw - 1.3, 0.66, ["Bot audit", [("RaidShield Pro", {"size": 11, "bold": False,
                                                                                           "color": MUTED})]],
         size=15, bold=True, color=WHITE)
    text(g2.shapes, x + 0.22, by + 1.25, cw - 0.44, 1.05,
         "Kai: real Snyk vulns, requests ADMINISTRATOR, evals user input. Tenki sandbox + Prelint confirm.",
         size=11, color=INK)
    # ask haven
    g3 = group(s.shapes)
    x = cards_x + 2 * (cw + 0.25)
    rect(g3.shapes, x, by, cw, bh, fill="3B82F6", alpha=0.07, line="3B82F6", line_alpha=0.35, radius=0.16)
    text(g3.shapes, x + 0.22, by + 0.25, cw - 0.44, 0.36, "Ask Haven", size=15, bold=True, color=WHITE)
    text(g3.shapes, x + 0.22, by + 0.6, cw - 0.44, 0.5, "\u201cWhich kind of incident is most common?\u201d",
         size=11, color=INK, italic=True)
    bars = [(0.95, ROSE, "phishing"), (0.62, AMBER, "spam"), (0.4, VIOLET_L, "raids"), (0.22, CYAN, "crisis")]
    bw = (cw - 0.44 - 3 * 0.1) / 4
    base = by + bh - 0.42
    for k, (v, col, lb_) in enumerate(bars):
        bx = x + 0.22 + k * (bw + 0.1)
        hh = 0.6 * v
        rect(g3.shapes, bx, base - hh, bw, hh, fill=col, alpha=0.85, line=None, radius=0.04)
        text(g3.shapes, bx - 0.05, base + 0.05, bw + 0.1, 0.25, lb_, size=8.5, color=MUTED, align="c")
    text(g3.shapes, x + 0.22, by + 1.07, cw - 0.44, 0.25, "PromptQL \u2192 SQL \u2192 chart", size=9.5,
         color="60A5FA", bold=True)
    fx += [(g1, "rise", 2900), (g2, "rise", 3050), (g3, "rise", 3200)]
    add_animations(s, fx)
    s.notes_slide.notes_text_frame.text = (
        "Here's the live demo. A scammer posts a fake free Premium link. In 0.4 seconds Nova flags it and "
        "quarantines it. Aria delegates to Rex, who detonates the link in a Celesto sandbox and captures the "
        "credential harvester at four seconds. At six seconds Rex warns the channel and Juno blocklists the "
        "domain. I approve the ban in one click. Same speed for crisis care, bot audits and data questions.")
    return s


def slide7(prs):
    s, ts = new_slide(prs, 7, "Trust by design", "Safety and trust are built into the company")
    fx = [(ts[0], "fade", 0), (ts[1], "rise", 150)]
    fr = screenshot_frame(s.shapes, "audit", 0.6, 1.85, 5.9, 4.7, "haven / audits / raidshield-pro")
    fx.append((fr, "zoom", 400, 800))
    items = [("user", EMERALD, "Human in the loop", "Bans need founder approval, granted in one click."),
             ("lifebuoy", "EC4899", "Crisis protocol", "Never auto-deleted; Mira replies and always pages a human."),
             ("eye", CYAN, "PII redaction", "Personal data is stripped before any model call through Flatkey."),
             ("doc", VIOLET_L, "Policy as specs", "Trust & safety and bot policies are specs checked by Prelint."),
             ("lock", AMBER, "Secure supply chain", "Snyk-gated CI on Tenki runners; bots audited before install.")]
    for i, (ic, col, h, d) in enumerate(items):
        y = 1.85 + i * 0.97
        g = group(s.shapes)
        rect(g.shapes, 6.85, y, 5.88, 0.82, radius=0.16)
        icon_circle(g.shapes, 7.03, y + 0.14, 0.54, col, ic)
        text(g.shapes, 7.78, y + 0.1, 4.8, 0.3, h, size=15, bold=True, color=WHITE)
        text(g.shapes, 7.78, y + 0.42, 4.85, 0.32, d, size=12, color=MUTED)
        fx.append((g, "left", 700 + i * 180))
    add_animations(s, fx)
    s.notes_slide.notes_text_frame.text = (
        "Safety has to be designed in. AI never bans on its own: every ban needs my one-click approval. "
        "Crisis messages are never auto-deleted and always reach a human. We redact personal data before any "
        "model call. Our policies are specs checked by Prelint, and our CI is gated by Snyk on Tenki runners. "
        "On the left, Kai's audit of a malicious bot, graded F.")
    return s


SPONSORS = [
    ("Kylon", "Workforce manifest for the 8 AI employees", VIOLET_L),
    ("Masky.AI", "Real faces and voices for Mira, Rex and Aria (+Maskord)", "EC4899"),
    ("BAND", "Agent-to-agent message bus; Aria routes over it", VIOLET_L),
    ("n8n", "Onboarding, human paging, enforcement logs, reports", AMBER),
    ("RocketRide", "Triage pipeline with a sub-millisecond fast path", CYAN),
    ("AdaL", "Browser workflow that drives link detonation", ROSE),
    ("Celesto", "Isolated computer where links are detonated", ROSE),
    ("Flatkey", "Tiered model gateway, failover, per-call cost", CYAN),
    ("PromptQL", "Natural-language questions over live safety data", "3B82F6"),
    ("Paritok", "Context compression for long incident threads", "3B82F6"),
    ("Finch", "Reusable skills as SKILL.md files", EMERALD),
    ("Tenki", "Sandbox test runs and CI runners", AMBER),
    ("Prelint", "Checks policies and specs before they ship", EMERALD),
    ("Snyk", "Dependency vulnerability scans for bots and CI", VIOLET_L),
]


def slide8(prs):
    s, ts = new_slide(prs, 8, "Sponsor stack", "14 sponsor tools, one AI-native company")
    fx = [(ts[0], "fade", 0), (ts[1], "rise", 150)]
    cols, gap = 5, 0.2
    tw = (12.13 - (cols - 1) * gap) / cols
    th = 1.45
    for i, (nm, d, col) in enumerate(SPONSORS):
        r, c = divmod(i, cols)
        x, y = 0.6 + c * (tw + gap), 1.85 + r * (th + gap)
        g = group(s.shapes)
        rect(g.shapes, x, y, tw, th, radius=0.16)
        rect(g.shapes, x + 0.22, y + 0.28, 0.14, 0.14, fill=col, alpha=1, line=None, radius=None,
             shape=MSO_SHAPE.OVAL, glow=(col, 6, 0.5))
        text(g.shapes, x + 0.48, y + 0.18, tw - 0.6, 0.34, nm, size=16, bold=True, color=WHITE, anchor="m")
        text(g.shapes, x + 0.22, y + 0.62, tw - 0.4, 0.75, d, size=12, color=MUTED)
        fx.append((g, "zoom", 300 + i * 90, 450))
    r, c = divmod(len(SPONSORS), cols)
    x, y = 0.6 + c * (tw + gap), 1.85 + r * (th + gap)
    g = group(s.shapes)
    rect(g.shapes, x, y, tw, th, grad=("5B21B6", "0E7490", 0), line=VIOLET_L, line_alpha=0.6, radius=0.16,
         glow=(VIOLET, 12, 0.35))
    text(g.shapes, x + 0.22, y + 0.2, tw - 0.44, 0.55, "14 \u2192 1", size=28, bold=True, color=WHITE)
    text(g.shapes, x + 0.22, y + 0.8, tw - 0.44, 0.55, "tools composed into one working company", size=12,
         color="E0E7FF")
    fx.append((g, "zoom", 300 + len(SPONSORS) * 90 + 100, 600))
    add_animations(s, fx)
    s.notes_slide.notes_text_frame.text = (
        "Haven is built on fourteen sponsor tools, and each one has a real job. Kylon defines the workforce, "
        "BAND carries agent messages, RocketRide does triage, AdaL and Celesto detonate links, Masky gives "
        "Mira a face and voice, Flatkey routes models, PromptQL answers questions, Paritok compresses context, "
        "Finch holds skills, n8n runs workflows, and Snyk, Tenki and Prelint keep us secure and on-policy.")
    return s


def slide9(prs):
    s, ts = new_slide(prs, 9, "Business", "Business model and traction")
    fx = [(ts[0], "fade", 0), (ts[1], "rise", 150)]
    tiers = [("Starter", "$49", "/mo", "Up to 1k members", INK, False),
             ("Growth", "$299", "/mo", "Up to 25k members", CYAN, True),
             ("Enterprise", "Custom", "", "Platforms & networks", VIOLET_L, False)]
    pw, pg, py, ph = 2.08, 0.2, 1.85, 2.45
    for i, (nm, price, per, d, col, hot) in enumerate(tiers):
        x = 0.6 + i * (pw + pg)
        g = group(s.shapes)
        if hot:
            rect(g.shapes, x, py, pw, ph, fill=CYAN, alpha=0.08, line=CYAN, line_alpha=0.7, line_w=1.25,
                 radius=0.18, glow=(CYAN, 12, 0.25))
            pill(g.shapes, x + pw - 1.08, py + 0.21, 0.88, 0.28, "POPULAR", color="0B0D14", fill=CYAN, alpha=1,
                 line=None, size=9)
        else:
            rect(g.shapes, x, py, pw, ph, radius=0.18)
        text(g.shapes, x + 0.22, py + 0.22, pw - 0.44, 0.3, nm, size=14, bold=True, color=col)
        text(g.shapes, x + 0.22, py + 0.7, pw - 0.3, 0.7, [[(price, {"size": 32 if price != "Custom" else 26,
                                                                    "bold": True, "color": WHITE}),
                                                           (per, {"size": 13, "color": MUTED})]], anchor="m")
        text(g.shapes, x + 0.22, py + 1.55, pw - 0.44, 0.7, d, size=12, color=MUTED)
        fx.append((g, "rise", 350 + i * 150))
    # unit economics
    ux, uw = 7.45, 5.28
    gu = group(s.shapes)
    rect(gu.shapes, ux, py, uw, ph, radius=0.18)
    text(gu.shapes, ux + 0.3, py + 0.22, uw - 0.6, 0.3, "COST PER INCIDENT", size=11, bold=True, color=CYAN, spacing=200)
    text(gu.shapes, ux + 0.3, py + 0.62, 2.2, 0.3, "Human moderator", size=13, color=INK)
    text(gu.shapes, ux + uw - 1.3, py + 0.62, 1.0, 0.3, "$4.20", size=15, bold=True, color=ROSE, align="r")
    rect(gu.shapes, ux + 0.3, py + 0.98, uw - 0.6, 0.26, fill=ROSE, alpha=0.85, line=None, radius=0.13)
    text(gu.shapes, ux + 0.3, py + 1.4, 2.2, 0.3, "Haven", size=13, color=INK)
    text(gu.shapes, ux + uw - 1.3, py + 1.4, 1.0, 0.3, "~$0.01", size=15, bold=True, color=EMERALD, align="r")
    rect(gu.shapes, ux + 0.3, py + 1.76, 0.12, 0.26, fill=EMERALD, alpha=1, line=None, radius=0.06)
    text(gu.shapes, ux + 0.55, py + 1.74, uw - 0.9, 0.3, "~400\u00d7 cheaper, on duty 24/7", size=12, color=MUTED,
         anchor="m")
    fx.append((gu, "left", 800))
    # bottom stats
    sy = 4.55
    stats = [("$15B+", "trust & safety market", VIOLET_L), ("30M+", "communities to serve", CYAN),
             ("40k", "members protected across 4 demo communities", EMERALD)]
    sw_ = 2.08
    for i, (num, lab, col) in enumerate(stats):
        x = 0.6 + i * (sw_ + 0.2)
        g = group(s.shapes)
        rect(g.shapes, x, sy, sw_, 2.1, fill=col, alpha=0.06, line=col, line_alpha=0.3, radius=0.18)
        text(g.shapes, x + 0.22, sy + 0.25, sw_ - 0.4, 0.65, num, size=30, bold=True, color=col)
        text(g.shapes, x + 0.22, sy + 1.0, sw_ - 0.44, 0.9, lab, size=12, color=MUTED)
        fx.append((g, "rise", 1100 + i * 150))
    # MRR chart (native)
    gch = rect(s.shapes, ux, sy, uw, 2.1, radius=0.18)
    lab = text(s.shapes, ux + 0.3, sy + 0.16, uw - 0.6, 0.3, "MRR PROJECTION (ILLUSTRATIVE)", size=11, bold=True,
               color=CYAN, spacing=200)
    cd = CategoryChartData()
    cd.categories = ["M1", "M3", "M6", "M9", "M12"]
    cd.add_series("MRR ($k)", (0.6, 3, 12, 30, 60))
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(ux + 0.15), Inches(sy + 0.42), Inches(uw - 0.3),
                            Inches(1.6), cd)
    gf.name = _uname("chart")
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(10)
    ch.font.color.rgb = hexrgb(MUTED)
    ch.font.name = FONT
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    ca = ch.category_axis
    ca.format.line.color.rgb = hexrgb("2A2E3E")
    ca.has_major_gridlines = False
    ca.tick_labels.font.color.rgb = hexrgb(MUTED)
    ser = ch.plots[0].series[0]
    ser.format.line.color.rgb = hexrgb(CYAN)
    ser.format.line.width = Pt(2.75)
    ser.smooth = True
    ser.marker.size = 7
    ser.marker.format.fill.solid()
    ser.marker.format.fill.fore_color.rgb = hexrgb(VIOLET_L)
    ser.marker.format.line.color.rgb = hexrgb(VIOLET_L)
    pl = ch.plots[0]
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.number_format = '"$"0.#"k"'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.ABOVE
    dl.font.size = Pt(10)
    dl.font.bold = True
    dl.font.color.rgb = hexrgb(INK)
    fx += [(gch, "fade", 1500), (lab, "fade", 1550), (gf, "fade", 1650, 800)]
    add_animations(s, fx)
    s.notes_slide.notes_text_frame.text = (
        "The business is simple SaaS. Starter is forty-nine dollars a month for up to a thousand members, "
        "Growth is two ninety-nine for up to twenty-five thousand, and Enterprise is custom for platforms. "
        "A human moderator costs about four twenty per incident; Haven costs about a cent. The market is over "
        "fifteen billion dollars, and today we protect four demo communities with forty thousand members.")
    return s


def slide10(prs):
    s, _ = new_slide(prs, 10)
    fx = []
    logo = picture(s.shapes, A["logo"], SW / 2 - 0.45, 0.65, 0.9, 0.9, name="!!herologo")
    fx.append((logo, "zoom", 0, 700))
    k = text(s.shapes, 1, 1.75, SW - 2, 0.3, "VISION", size=12, bold=True, color=CYAN, spacing=300, align="c")
    fx.append((k, "fade", 200))
    v = text(s.shapes, 1, 2.1, SW - 2, 1.3,
             ["Every community gets a", [("24/7 AI Trust & Safety team", {"grad": (VIOLET_L, CYAN)})]],
             size=38, bold=True, color=WHITE, align="c", anchor="m")
    fx.append((v, "rise", 350, 800))
    asks = [("Design partners", "10 communities to pilot Starter and Growth this quarter", CYAN),
            ("Platform partners", "Discord, Slack and Telegram app-directory distribution", VIOLET_L),
            ("Support", "Sponsor credits and prize funding to scale the AI team", EMERALD)]
    cw, gap, y = 3.75, 0.3, 3.75
    x0 = SW / 2 - (3 * cw + 2 * gap) / 2
    lab = text(s.shapes, x0, y - 0.05, 6, 0.3, "OUR ASK", size=11, bold=True, color=MUTED, spacing=200)
    fx.append((lab, "fade", 800))
    for i, (h, d, col) in enumerate(asks):
        x = x0 + i * (cw + gap)
        g = group(s.shapes)
        rect(g.shapes, x, y + 0.35, cw, 1.4, fill=col, alpha=0.07, line=col, line_alpha=0.4, radius=0.18)
        text(g.shapes, x + 0.3, y + 0.55, cw - 0.6, 0.35, h, size=17, bold=True, color=WHITE)
        text(g.shapes, x + 0.3, y + 0.95, cw - 0.6, 0.7, d, size=13, color=MUTED)
        fx.append((g, "rise", 950 + i * 180))
    cta = group(s.shapes)
    rect(cta.shapes, SW / 2 - 2.6, 5.9, 5.2, 0.62, grad=(VIOLET, "0891B2", 0), line=None, radius=0.31,
         glow=(VIOLET, 16, 0.4))
    text(cta.shapes, SW / 2 - 2.6, 5.9, 5.2, 0.62, "github.com/vnmoorthy/haven  \u2192", size=17, bold=True,
         color=WHITE, align="c", anchor="m", link=GITHUB)
    fx.append((cta, "zoom", 1600, 700))
    f = text(s.shapes, 1, 6.75, SW - 2, 0.3, "Built by Moorthy + 8 AI employees  ·  Thank you",
             size=12, color=MUTED, align="c")
    fx.append((f, "fade", 2100))
    add_animations(s, fx)
    s.notes_slide.notes_text_frame.text = (
        "Our vision: every online community, big or small, gets its own 24/7 AI Trust and Safety team. "
        "We're looking for design-partner communities to pilot Haven, platform partners for distribution, "
        "and sponsor support to scale the AI team. Everything is open on GitHub at vnmoorthy/haven. "
        "Thank you. Haven, built by one human and eight AI employees.")
    return s


def main():
    global A
    A = build_assets()
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    prs.core_properties.title = "Haven — Pitch"
    prs.core_properties.author = "Moorthy"
    prs.core_properties.subject = "The AI Trust & Safety team for every online community"
    for fn in (slide1, slide2, slide3, slide4, slide5, slide6, slide7, slide8, slide9, slide10):
        s = fn(prs)
        # transition must precede timing in the slide XML
        timing = s._element.find(qn("p:timing"))
        if timing is not None:
            s._element.remove(timing)
        add_morph_transition(s)
        if timing is not None:
            s._element.append(timing)
    prs.save(OUT)
    shots = [k for k in ("mission-control", "incident-room", "audit") if os.path.exists(os.path.join(SHOTS, k + ".png"))]
    print("Wrote %s (%d slides). Real screenshots used: %s" % (OUT, len(prs.slides), ", ".join(shots) or "none (mocks)"))


if __name__ == "__main__":
    main()
